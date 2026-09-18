import csv

import pytest

from ksa_pipeline import export, fingerprint, paths, pii, scoring
from test_scoring import merchant, setup as scoring_setup, write
from ksa_pipeline.resolve import MERCHANT_COLUMNS


def setup(tmp_path, monkeypatch):
    scoring_setup(tmp_path, monkeypatch)
    monkeypatch.setattr(paths, "ROOT", tmp_path)
    monkeypatch.setattr(paths, "DATA", tmp_path)
    m = fingerprint.read_csv(tmp_path / "interim" / "merchants.csv")
    for r in m:
        if r["merchant_id"] in ("a1", "b1"):
            r.update(phone_e164="+966550000123", phone_type="mobile", website_domain="clinic.sa", instagram_handle="clinic")
    write(tmp_path / "interim" / "merchants.csv", m, MERCHANT_COLUMNS)
    write(tmp_path / "interim" / "enrich.csv",
          [{"merchant_id": "a1", "owner_name_in_source": "True", "input_kind": "page_text",
            "owner_name_inferred": "د. سارة", "hook_ar_inferred": "مرحبا"}],
          ["merchant_id", "owner_name_in_source", "input_kind", "owner_name_inferred", "hook_ar_inferred"])
    (tmp_path / "sources.yaml").write_text(
        "sources:\n  - id: google_maps\n    role: primary\n    status: accepted\n    decided_on: 2026-09-16\n    note: test\n", encoding="utf-8")
    scoring.run()


def test_masked_export_has_no_numbers_and_keeps_the_top_list(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    out = export.run()
    masked = list(csv.DictReader((tmp_path / "output" / "public" / "scored_masked.csv").open(encoding="utf-8")))
    assert out["merchants"] == len(masked) and out["top"] >= 1
    text = (tmp_path / "output" / "public" / "scored_masked.csv").read_text(encoding="utf-8")
    assert not pii.find_mobiles(text) and "+9665\u2022\u2022\u2022\u2022\u2022123" in text
    assert "owner_name_inferred" not in text                                   # names stay out of the masked export
    top = list(csv.DictReader((tmp_path / "output" / "public" / "top50_masked.csv").open(encoding="utf-8")))
    assert top and top[0]["top_rank"] == "1" and "hook_ar_ai_drafted" in top[0]
    assert top[0]["contact_note"] == "mobile in Google Maps"
    assert (tmp_path / "output" / "public" / "ksa_sme_pipeline.xlsx").exists()
    assert not (tmp_path / "private" / "top50_contacts.csv").exists()          # private file only on demand


def test_private_export_is_written_only_with_the_flag(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    export.run(private=True)
    rows = list(csv.DictReader((tmp_path / "private" / "top50_contacts.csv").open(encoding="utf-8")))
    assert rows[0]["phone_e164"] == "+966550000123" and rows[0]["owner_name_inferred"] == "د. سارة"


def test_export_columns_follow_the_schema():
    from ksa_pipeline.schema import export_columns
    assert export.TOP_COLUMNS == [c["name"] for c in export_columns()]
    assert "hook_ar_ai_drafted" not in export.MASKED_COLUMNS and "phone_e164" not in export.TOP_COLUMNS


def test_check_masked_catches_a_leak():
    with pytest.raises(AssertionError):
        export.check_masked([{"merchant_id": "x", "note": "call 0550000123"}])


def test_workbook_sheets(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    export.run()
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "output" / "public" / "ksa_sme_pipeline.xlsx")
    assert wb.sheetnames == ["README", "Top-50", "Reserve", "Scored", "Sources"]
    assert wb["Top-50"].freeze_panes == "A2" and wb["README"]["A1"].value.startswith("KSA SME")


def test_pii_scan_reads_the_workbook(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    export.run()
    book = tmp_path / "output" / "public" / "ksa_sme_pipeline.xlsx"
    assert pii.scan([book]) == {}
    from openpyxl import load_workbook
    wb = load_workbook(book)
    wb["Scored"]["Z1"] = "0550000123"
    wb.save(book)
    assert pii.scan([book]) == {str(book): 1}
