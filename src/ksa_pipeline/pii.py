"""Guard: no unmasked Saudi mobile numbers in committed files (wired into pre-commit)."""
from __future__ import annotations

import re
from pathlib import Path

from .normalize import to_ascii_digits

# Separators inside one phone number: spaces, tabs, dashes, dots, brackets — never a line break,
# otherwise a digit at the end of one CSV row glues onto the date that starts the next row.
_JOIN_SEPARATORS = re.compile(r"(?<=\d)[ \t\-.()]+(?=\d)")
_ISO_DATE = re.compile(r"(?<!\d)\d{4}-\d{2}-\d{2}(?!\d)")
_KSA_MOBILE = re.compile(r"(?<!\d)(?:\+?966|00966|0)?5\d{8}(?!\d)")


def find_mobiles(text: str) -> list[str]:
    text = _ISO_DATE.sub(" DATE ", to_ascii_digits(text))
    return _KSA_MOBILE.findall(_JOIN_SEPARATORS.sub("", text))


def mask_mobile(number: str) -> str:
    """+966550000123 -> +9665•••••123 (keeps the country code and the last three digits)."""
    digits = re.sub(r"\D", "", to_ascii_digits(number or ""))
    if not digits:
        return ""
    if len(digits) <= 7:
        return "\u2022" * len(digits)
    prefix = "+" if digits.startswith("966") else ""
    return prefix + digits[:4] + "\u2022" * (len(digits) - 7) + digits[-3:]


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)


def scan(paths: list[Path]) -> dict[str, int]:
    hits = {}
    for p in paths:
        if p.suffix.lower() == ".xlsx":                      # committed exports are workbooks too
            n = len(find_mobiles(_xlsx_text(p)))
            if n:
                hits[str(p)] = n
            continue
        if p.suffix.lower() not in {".csv", ".md", ".yaml", ".yml", ".json", ".txt"}:
            continue
        n = len(find_mobiles(p.read_text(encoding="utf-8", errors="ignore")))
        if n:
            hits[str(p)] = n
    return hits
